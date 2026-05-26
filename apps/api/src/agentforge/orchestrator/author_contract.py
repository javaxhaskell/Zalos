"""Author upload profiling for LLM-first contract planning.

This module intentionally does not infer workflow contracts, formulas,
exception rules, summaries, or finance-specific business semantics. Its job is
to describe the uploaded tabular file for the model.
"""

from __future__ import annotations

import csv
from decimal import Decimal, InvalidOperation
from pathlib import Path

from agentforge.schemas.author_output_contract import (
    ColumnNullProfile,
    ColumnUniquenessProfile,
    TabularColumnProfile,
    TabularFileProfile,
    WorkbookSheetProfile,
)

CONTRACT_REL = "generated/author_output_contract.json"
NORMALISED_INPUT_REL = "uploads/normalised_input.csv"

_ID_COLUMN_HINTS = frozenset(
    {
        "id",
        "row_id",
        "txn_id",
        "transaction_id",
        "payment_id",
        "charge_id",
        "event_id",
        "invoice_id",
        "expense_id",
        "merchant_id",
        "settlement_id",
        "reference",
        "payout_id",
    }
)
_DATE_COLUMN_HINTS = frozenset(
    {
        "date",
        "txn_date",
        "transaction_date",
        "settlement_date",
        "invoice_date",
        "due_date",
        "payment_date",
        "posted_date",
        "expense_date",
    }
)
_AMOUNT_COLUMN_HINTS = frozenset(
    {
        "amount",
        "gross",
        "gross_amount",
        "fee",
        "fee_amount",
        "net",
        "net_amount",
        "payment_amount",
        "debit",
        "credit",
        "policy_limit",
    }
)
_STATUS_COLUMN_HINTS = frozenset(
    {"status", "state", "approval_status", "reconciliation_status"}
)
_CATEGORY_COLUMN_HINTS = frozenset({"category", "type", "expense_category"})
_GROUP_COLUMN_HINTS = frozenset(
    {
        "account",
        "account_code",
        "aging_bucket",
        "batch_id",
        "category",
        "currency",
        "customer",
        "customer_id",
        "employee_id",
        "gl_account",
        "merchant",
        "processor",
        "settlement_batch",
        "vendor",
        "vendor_id",
    }
)
_REFERENCE_COLUMN_HINTS = frozenset(
    {
        "counterparty",
        "customer_reference",
        "description",
        "memo",
        "merchant",
        "notes",
        "reference",
        "remittance_id",
    }
)


def profile_tabular_file(
    *,
    workspace: Path,
    upload_path: Path,
    upload_format: str,
) -> TabularFileProfile:
    """Inspect CSV/XLSX shape without choosing workflow semantics."""

    rel_original = str(upload_path.relative_to(workspace)).replace("\\", "/")
    if upload_format == "csv":
        columns = _read_csv_header(upload_path)
        rows = _read_csv_rows(upload_path, limit=None)
        sample_rows = rows[:5]
        return TabularFileProfile(
            original_upload_path=rel_original,
            input_format="csv",
            selected_sheet=None,
            normalized_input_path=rel_original,
            agent_input_path=rel_original,
            columns=columns,
            row_count=len(rows),
            column_profiles=_profile_columns(columns, sample_rows),
            candidate_id_columns=_candidate_columns(columns, _ID_COLUMN_HINTS),
            candidate_date_columns=_candidate_columns(columns, _DATE_COLUMN_HINTS),
            candidate_amount_columns=_candidate_columns(columns, _AMOUNT_COLUMN_HINTS),
            candidate_status_columns=_candidate_columns(columns, _STATUS_COLUMN_HINTS),
            candidate_category_columns=_candidate_columns(columns, _CATEGORY_COLUMN_HINTS),
            candidate_grouping_columns=_candidate_columns(columns, _GROUP_COLUMN_HINTS),
            candidate_reference_columns=_candidate_columns(columns, _REFERENCE_COLUMN_HINTS),
            null_profiles=_null_profiles(columns, rows),
            uniqueness_profiles=_uniqueness_profiles(columns, rows),
            sample_rows=sample_rows,
            plausible_sheets=[],
            workbook_sheets=[],
        )

    if upload_format != "xlsx":
        raise ValueError(f"unsupported tabular upload format: {upload_format}")

    workbook_profile = _profile_xlsx_workbook(upload_path)
    if workbook_profile["selected_sheet"] is None:
        raise ValueError("workbook has no readable sheet with a header row")

    normalized_abs = workspace / NORMALISED_INPUT_REL
    normalized_abs.parent.mkdir(parents=True, exist_ok=True)
    rows_for_csv: list[list[str]] = workbook_profile["rows_for_csv"]
    columns = list(workbook_profile["columns"])
    with normalized_abs.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(columns)
        writer.writerows(rows_for_csv)

    sample_rows = list(workbook_profile["sample_rows"])
    all_rows = list(workbook_profile["all_rows"])
    return TabularFileProfile(
        original_upload_path=rel_original,
        input_format="xlsx",
        selected_sheet=workbook_profile["selected_sheet"],
        normalized_input_path=NORMALISED_INPUT_REL,
        agent_input_path=NORMALISED_INPUT_REL,
        columns=columns,
        row_count=len(all_rows),
        column_profiles=_profile_columns(columns, sample_rows),
        candidate_id_columns=_candidate_columns(columns, _ID_COLUMN_HINTS),
        candidate_date_columns=_candidate_columns(columns, _DATE_COLUMN_HINTS),
        candidate_amount_columns=_candidate_columns(columns, _AMOUNT_COLUMN_HINTS),
        candidate_status_columns=_candidate_columns(columns, _STATUS_COLUMN_HINTS),
        candidate_category_columns=_candidate_columns(columns, _CATEGORY_COLUMN_HINTS),
        candidate_grouping_columns=_candidate_columns(columns, _GROUP_COLUMN_HINTS),
        candidate_reference_columns=_candidate_columns(columns, _REFERENCE_COLUMN_HINTS),
        null_profiles=_null_profiles(columns, all_rows),
        uniqueness_profiles=_uniqueness_profiles(columns, all_rows),
        sample_rows=sample_rows,
        plausible_sheets=list(workbook_profile["sheet_names"]),
        workbook_sheets=list(workbook_profile["sheet_profiles"]),
        primary_sheet=workbook_profile["selected_sheet"],
    )


def _read_csv_header(path: Path) -> list[str]:
    with path.open(newline="", encoding="utf-8", errors="replace") as handle:
        return next(csv.reader(handle), [])


def _read_csv_rows(path: Path, *, limit: int | None) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8", errors="replace") as handle:
        reader = csv.DictReader(handle)
        rows = list(reader)
    if limit is None:
        return rows
    return rows[:limit]


def _profile_xlsx_workbook(upload_path: Path) -> dict[str, object]:
    import openpyxl

    workbook = openpyxl.load_workbook(upload_path, read_only=True, data_only=True)
    try:
        sheet_names = list(workbook.sheetnames)
        sheet_profiles: list[WorkbookSheetProfile] = []
        selected_sheet: str | None = None
        selected_columns: list[str] = []
        selected_rows: list[dict[str, str]] = []
        rows_for_csv: list[list[str]] = []

        for sheet_name in sheet_names:
            sheet = workbook[sheet_name]
            rows_iter = sheet.iter_rows(values_only=True)
            header_row = next(rows_iter, None)
            if not header_row:
                sheet_profiles.append(
                    WorkbookSheetProfile(
                        sheet_name=sheet_name,
                        row_count=0,
                        columns=[],
                        sample_rows=[],
                    )
                )
                continue

            columns = [
                str(cell).strip()
                for cell in header_row
                if cell is not None and str(cell).strip()
            ]
            data_rows: list[list[str]] = []
            sample_rows: list[dict[str, str]] = []
            all_rows: list[dict[str, str]] = []
            for raw_row in rows_iter:
                values = ["" if cell is None else str(cell) for cell in raw_row]
                if not any(value.strip() for value in values):
                    continue
                trimmed = values[: len(columns)]
                if len(trimmed) < len(columns):
                    trimmed.extend([""] * (len(columns) - len(trimmed)))
                row_dict = {columns[i]: trimmed[i] for i in range(len(columns))}
                data_rows.append(trimmed)
                all_rows.append(row_dict)
                if len(sample_rows) < 5:
                    sample_rows.append(row_dict)

            sheet_profiles.append(
                WorkbookSheetProfile(
                    sheet_name=sheet_name,
                    row_count=len(all_rows),
                    columns=columns,
                    sample_rows=sample_rows,
                    inferred_role="primary_data" if selected_sheet is None and columns else "unknown",
                )
            )
            if selected_sheet is None and columns:
                selected_sheet = sheet_name
                selected_columns = columns
                selected_rows = all_rows
                rows_for_csv = data_rows
    finally:
        workbook.close()

    return {
        "sheet_names": sheet_names,
        "sheet_profiles": sheet_profiles,
        "selected_sheet": selected_sheet,
        "columns": selected_columns,
        "sample_rows": selected_rows[:5],
        "all_rows": selected_rows,
        "rows_for_csv": rows_for_csv,
    }


def _profile_columns(
    columns: list[str],
    rows: list[dict[str, str]],
) -> list[TabularColumnProfile]:
    profiles: list[TabularColumnProfile] = []
    for column in columns:
        samples = [(row.get(column) or "").strip() for row in rows[:5]]
        non_empty = [sample for sample in samples if sample]
        profiles.append(
            TabularColumnProfile(
                name=column,
                inferred_type=_infer_column_type(non_empty),
                sample_values=non_empty[:3],
            )
        )
    return profiles


def _infer_column_type(values: list[str]) -> str:
    if not values:
        return "string"
    numeric = 0
    for value in values:
        text = value.strip()
        if not text:
            continue
        try:
            Decimal(text.replace(",", ""))
            numeric += 1
        except InvalidOperation:
            return "string"
    return "decimal" if numeric else "string"


def _candidate_columns(columns: list[str], hints: frozenset[str]) -> list[str]:
    mapping = {column.strip().lower(): column for column in columns if column.strip()}
    return [mapping[hint] for hint in sorted(hints) if hint in mapping]


def _null_profiles(
    columns: list[str],
    rows: list[dict[str, str]],
) -> list[ColumnNullProfile]:
    profiles: list[ColumnNullProfile] = []
    for column in columns:
        empty = 0
        non_empty = 0
        for row in rows:
            if (row.get(column) or "").strip():
                non_empty += 1
            else:
                empty += 1
        profiles.append(
            ColumnNullProfile(
                column=column,
                empty_count=empty,
                non_empty_count=non_empty,
            )
        )
    return profiles


def _uniqueness_profiles(
    columns: list[str],
    rows: list[dict[str, str]],
) -> list[ColumnUniquenessProfile]:
    profiles: list[ColumnUniquenessProfile] = []
    for column in columns:
        values = [(row.get(column) or "").strip() for row in rows if (row.get(column) or "").strip()]
        unique_values = len(set(values))
        profiles.append(
            ColumnUniquenessProfile(
                column=column,
                unique_values=unique_values,
                total_values=len(values),
                is_unique=bool(values) and unique_values == len(values),
            )
        )
    return profiles


__all__ = [
    "CONTRACT_REL",
    "NORMALISED_INPUT_REL",
    "profile_tabular_file",
]
